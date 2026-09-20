# backend/app/modules/reportes/exportacion.py
# CU20 - Exportacion de reportes a PDF y Excel (READ-ONLY).
#
# Recibe el DTO ya calculado por service.py (el MISMO que se devuelve como JSON)
# y lo convierte a archivo. NO consulta la base de datos ni recalcula nada: por
# eso el archivo coincide exactamente con lo que el usuario ve en pantalla.
#
# Arquitectura (sin duplicar logica entre formatos):
#   DTO  --(un constructor por reporte)-->  Documento  --> renderizador xlsx | pdf
# `Documento` es un modelo neutro (titulo, filtros aplicados, indicadores,
# secciones tabulares, notas). Cada reporte se describe UNA vez; Excel y PDF
# solo dibujan. Los caracteres no latinos se sustituyen en el PDF (fuentes base
# de reportlab); el Excel conserva el texto tal cual.
#
# Seguridad:
# - Excel: todo texto se escribe como cadena (data_type='s'): un producto o
#   motivo que empiece con "=", "+", "-" o "@" NO se evalua como formula.
# - PDF: el texto se escapa antes de pasarlo a Paragraph (sin marcado inyectado).
# Fechas/horas: UTC (misma estrategia del resto de CU20).
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal, Optional
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.reporte import (
    MENSAJE_SIN_DATOS,
    DevolucionesReporteResponse,
    InventarioSituacionResponse,
    ProductosMasVendidosResponse,
    RendimientoVendedoresResponse,
    VentasPeriodoResponse,
)

Formato = Literal["pdf", "xlsx"]
TipoReporte = Literal[
    "ventas", "productos-mas-vendidos", "inventario", "devoluciones", "rendimiento-vendedores"
]

MEDIA_TYPE: dict[str, str] = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Compresion de las paginas del PDF. Las pruebas la apagan para poder buscar
# texto dentro del binario; en produccion queda activa.
COMPRIMIR_PDF = True

COLOR_PRIMARIO = "1E4D8C"  # azul Attention
COLOR_SUAVE = "F0F7FF"

_CANAL = {"ONLINE": "Online", "POS": "POS (mostrador)"}


# ---------------------------------------------------------------------------
# Modelo neutro del documento
# ---------------------------------------------------------------------------
TipoColumna = Literal["texto", "entero", "moneda", "decimal"]


@dataclass
class Columna:
    titulo: str
    tipo: TipoColumna = "texto"


@dataclass
class Seccion:
    titulo: str
    columnas: list[Columna]
    filas: list[list[Any]]
    nota: Optional[str] = None
    hoja: Optional[str] = None  # nombre corto de la hoja Excel (max. 31 caracteres)


@dataclass
class Indicador:
    etiqueta: str
    valor: Any
    tipo: TipoColumna = "entero"
    nota: Optional[str] = None


@dataclass
class Documento:
    tipo: str
    titulo: str
    periodo: tuple[str, str]
    filtros: list[tuple[str, str]]
    sin_datos: bool
    mensaje: Optional[str]
    indicadores: list[Indicador] = field(default_factory=list)
    secciones: list[Seccion] = field(default_factory=list)
    notas: list[str] = field(default_factory=list)


@dataclass
class ArchivoExportado:
    contenido: bytes
    media_type: str
    extension: str
    nombre: str


# ---------------------------------------------------------------------------
# Formato de valores (solo texto; ninguna operacion de negocio)
# ---------------------------------------------------------------------------
def _num(valor: Any) -> Optional[float]:
    return None if valor is None else float(valor)


def _texto_valor(valor: Any, tipo: TipoColumna) -> str:
    if valor is None:
        return "N/D"
    if tipo == "moneda":
        return f"Bs. {float(valor):,.2f}"
    if tipo == "entero":
        return f"{int(valor):,}"
    if tipo == "decimal":
        return f"{float(valor):,.2f}"
    return str(valor)


def _fecha_hora(valor: Optional[datetime]) -> str:
    return valor.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M") if valor else "—"


def _filtros_comunes(
    d: Any, categoria: Optional[str], etiqueta_periodo: str = "Período (UTC)"
) -> list[tuple[str, str]]:
    f = d.filtros
    cat = categoria or ("Todas" if f.categoria_id is None else f"#{f.categoria_id}")
    return [
        (etiqueta_periodo, f"{f.fecha_inicio} a {f.fecha_fin}"),
        ("Categoría", cat),
        ("Canal de venta", _CANAL.get(f.canal_venta, "Todos")),
    ]


def _extra(extras: dict[str, Any], clave: str, etiqueta: str) -> list[tuple[str, str]]:
    valor = extras.get(clave)
    return [] if valor in (None, "") else [(etiqueta, str(valor))]


# ---------------------------------------------------------------------------
# Constructores del documento (uno por reporte; los datos vienen del DTO)
# ---------------------------------------------------------------------------
def _doc_ventas(d: VentasPeriodoResponse, categoria: Optional[str], extras: dict) -> Documento:
    doc = Documento(
        tipo="ventas",
        titulo="Reporte de ventas",
        periodo=(str(d.filtros.fecha_inicio), str(d.filtros.fecha_fin)),
        filtros=_filtros_comunes(d, categoria),
        sin_datos=d.sin_datos,
        mensaje=d.mensaje,
    )
    if d.sin_datos:
        return doc
    doc.indicadores = [
        Indicador("Ingresos por productos", d.ingresos_productos, "moneda", "Sin costo de envío"),
        Indicador("Nº de ventas", d.num_ventas, "entero", "Solo ventas pagadas"),
        Indicador("Unidades vendidas", d.unidades_vendidas, "entero"),
        Indicador("Ticket promedio", d.ticket_promedio, "moneda"),
    ]
    if d.total_facturado is not None:
        doc.indicadores.append(Indicador("Total facturado", d.total_facturado, "moneda", "Incluye envío"))
    if d.total_envios is not None:
        doc.indicadores.append(Indicador("Envíos cobrados", d.total_envios, "moneda"))
    else:
        doc.notas.append(
            "Con filtro de categoría no se informan el total facturado ni los envíos: "
            "el envío no es atribuible a una categoría."
        )
    doc.secciones = [
        Seccion(
            "Detalle por día (solo días con ventas)",
            [Columna("Fecha"), Columna("Ventas", "entero"), Columna("Unidades", "entero"), Columna("Ingresos", "moneda")],
            [[str(p.fecha), p.num_ventas, p.unidades, p.ingresos] for p in d.por_fecha if p.num_ventas > 0],
            hoja="Por día",
        ),
        Seccion(
            "Por categoría",
            [Columna("Categoría"), Columna("Unidades", "entero"), Columna("Ingresos", "moneda")],
            [[c.categoria, c.unidades, c.ingresos] for c in d.por_categoria],
        ),
        Seccion(
            "Por canal",
            [Columna("Canal"), Columna("Ventas", "entero"), Columna("Ingresos", "moneda")],
            [[_CANAL.get(c.canal, c.canal), c.num_ventas, c.ingresos] for c in d.por_canal],
        ),
        Seccion(
            "Por método de pago",
            [Columna("Método de pago"), Columna("Ventas", "entero"), Columna("Ingresos", "moneda")],
            [[m.metodo_pago, m.num_ventas, m.ingresos] for m in d.por_metodo_pago],
        ),
    ]
    return doc


def _doc_top(d: ProductosMasVendidosResponse, categoria: Optional[str], extras: dict) -> Documento:
    doc = Documento(
        tipo="productos-mas-vendidos",
        titulo="Reporte de productos más vendidos",
        periodo=(str(d.filtros.fecha_inicio), str(d.filtros.fecha_fin)),
        filtros=_filtros_comunes(d, categoria) + [("Top", str(d.top))],
        sin_datos=d.sin_datos,
        mensaje=d.mensaje,
    )
    if d.sin_datos:
        return doc
    doc.indicadores = [
        Indicador("Productos en el ranking", len(d.items), "entero", f"Top solicitado: {d.top}"),
        Indicador("Productos distintos vendidos", d.total_productos_vendidos, "entero", "En el período, antes del top"),
    ]
    doc.secciones = [
        Seccion(
            "Ranking por unidades vendidas",
            [Columna("#", "entero"), Columna("Producto"), Columna("Categoría"),
             Columna("Cantidad vendida", "entero"), Columna("Total generado", "moneda"), Columna("Nº de ventas", "entero")],
            [[i.posicion, i.producto, i.categoria, i.cantidad_vendida, i.total_generado, i.num_ventas] for i in d.items],
        )
    ]
    return doc


def _doc_inventario(d: InventarioSituacionResponse, categoria: Optional[str], extras: dict) -> Documento:
    doc = Documento(
        tipo="inventario",
        titulo="Reporte de inventario",
        periodo=(str(d.filtros.fecha_inicio), str(d.filtros.fecha_fin)),
        filtros=_filtros_comunes(d, categoria, "Ventana de rotación (UTC)")
        + _extra(extras, "nivel_stock", "Nivel de stock")
        + [("Máx. de filas", str(d.limite))],
        sin_datos=d.sin_datos,
        mensaje=d.mensaje,
    )
    if d.sin_datos:
        return doc
    doc.indicadores = [
        Indicador("Productos", d.total_productos, "entero", "Activos y agotados"),
        Indicador("Unidades en stock", d.stock_total_unidades, "entero", "Stock global"),
        Indicador("Valor del inventario", d.valor_inventario, "moneda", "A precio de venta"),
        Indicador("Agotados", d.agotados, "entero", "Stock = 0"),
    ]
    doc.secciones = [
        Seccion(
            "Productos por nivel de stock",
            [Columna("Nivel"), Columna("Productos", "entero")],
            [["Crítico (< 5)", d.por_nivel.get("CRITICO", 0)], ["Bajo (< 15)", d.por_nivel.get("BAJO", 0)], ["OK", d.por_nivel.get("OK", 0)]],
        ),
        Seccion(
            "Detalle de inventario",
            [Columna("Producto"), Columna("Categoría"), Columna("Estado"), Columna("Stock", "entero"),
             Columna("Nivel"), Columna("Precio", "moneda"), Columna("Valor en stock", "moneda"),
             Columna("Vendidas (período)", "entero"), Columna("Rotación (aprox.)", "decimal")],
            [[i.producto, i.categoria, i.estado, i.stock_actual, i.nivel_stock, i.precio_venta,
              i.valor_stock, i.unidades_vendidas, i.rotacion if i.rotacion_disponible else None]
             for i in d.items],
            nota=(f"Mostrando {len(d.items)} de {d.total_filas} productos (límite {d.limite}); "
                  "los indicadores corresponden al total." if len(d.items) < d.total_filas else None),
        ),
    ]
    doc.notas = [
        "Rotación aproximada = unidades vendidas del período ÷ stock actual. No es la rotación contable "
        "(no hay costo de compra) y no se calcula (N/D) cuando el stock es 0.",
        "El stock es el actual; las fechas y el canal solo afectan a las unidades vendidas y a la rotación.",
    ]
    return doc


def _doc_devoluciones(d: DevolucionesReporteResponse, categoria: Optional[str], extras: dict) -> Documento:
    doc = Documento(
        tipo="devoluciones",
        titulo="Reporte de devoluciones",
        periodo=(str(d.filtros.fecha_inicio), str(d.filtros.fecha_fin)),
        filtros=_filtros_comunes(d, categoria)
        + ([("Estado", d.estado)] if d.estado else [])
        + _extra(extras, "top", "Top de productos")
        + _extra(extras, "limite_detalle", "Máx. filas del detalle"),
        sin_datos=d.sin_datos,
        mensaje=d.mensaje,
    )
    if d.sin_datos:
        return doc
    doc.indicadores = [
        Indicador("Devoluciones", d.num_devoluciones, "entero", "Todos los estados"),
        Indicador("Unidades devueltas", d.unidades_devueltas, "entero", "Sin rechazadas"),
        Indicador("Importe devuelto", d.importe_total, "moneda", "Sin rechazadas"),
        Indicador("Importe completado", d.importe_completado, "moneda", "Solo completadas"),
    ]
    doc.secciones = [
        Seccion(
            "Por estado",
            [Columna("Estado"), Columna("Devoluciones", "entero"), Columna("Unidades", "entero"), Columna("Importe", "moneda")],
            [[e.estado, e.num_devoluciones, e.unidades, e.importe] for e in d.por_estado],
        ),
        Seccion(
            "Por día de solicitud (solo días con devoluciones)",
            [Columna("Fecha"), Columna("Devoluciones", "entero"), Columna("Unidades", "entero"), Columna("Importe", "moneda")],
            [[str(p.fecha), p.num_devoluciones, p.unidades, p.importe] for p in d.por_fecha if p.num_devoluciones > 0],
            hoja="Por día",
        ),
        Seccion(
            "Por producto",
            [Columna("Producto"), Columna("Categoría"), Columna("Unidades", "entero"), Columna("Importe", "moneda")],
            [[p.producto, p.categoria, p.unidades, p.importe] for p in d.por_producto],
        ),
        Seccion(
            "Por categoría",
            [Columna("Categoría"), Columna("Unidades", "entero"), Columna("Importe", "moneda")],
            [[c.categoria, c.unidades, c.importe] for c in d.por_categoria],
        ),
        Seccion(
            "Detalle de devoluciones",
            [Columna("#", "entero"), Columna("Solicitud (UTC)"), Columna("Venta"), Columna("Estado"),
             Columna("Motivo"), Columna("Unidades", "entero"), Columna("Importe solicitado", "moneda")],
            [[x.id_devolucion, _fecha_hora(x.fecha_solicitud), x.codigo_venta, x.estado, x.motivo, x.unidades, x.importe]
             for x in d.detalle],
        ),
    ]
    doc.notas = [
        "Las devoluciones rechazadas cuentan en el número de devoluciones y en el desglose por estado, "
        "pero no en unidades ni importes (no hubo reembolso). Las fechas corresponden a la solicitud (UTC)."
    ]
    return doc


def _doc_rendimiento(d: RendimientoVendedoresResponse, categoria: Optional[str], extras: dict) -> Documento:
    sin_datos = len(d.items) == 0
    doc = Documento(
        tipo="rendimiento-vendedores",
        titulo="Reporte de rendimiento de vendedores",
        periodo=(str(d.fecha_desde), str(d.fecha_hasta)),
        filtros=[("Período (UTC)", f"{d.fecha_desde} a {d.fecha_hasta}")]
        + [("Canal de venta", _CANAL.get(extras.get("tipo_venta") or "", "Todos"))],
        sin_datos=sin_datos,
        mensaje=MENSAJE_SIN_DATOS if sin_datos else None,
    )
    if sin_datos:
        return doc
    doc.indicadores = [
        Indicador("Vendedores con ventas", d.total_vendedores, "entero"),
        Indicador("Ingresos", d.total_ingresos, "moneda", "Suma de `total` de sus ventas"),
        Indicador("Operaciones", d.total_operaciones, "entero"),
    ]
    doc.secciones = [
        Seccion(
            "Rendimiento por vendedor",
            [Columna("Vendedor"), Columna("Correo"), Columna("Ventas", "entero"), Columna("Ingresos", "moneda"),
             Columna("Ticket promedio", "moneda"), Columna("Tipos"), Columna("Primera venta (UTC)"), Columna("Última venta (UTC)")],
            [[v.nombre, v.correo, v.total_ventas, v.total_ingresos, v.ticket_promedio,
              " · ".join(f"{k}: {n}" for k, n in v.tipos_venta.items()),
              _fecha_hora(v.primera_venta), _fecha_hora(v.ultima_venta)]
             for v in d.items],
        )
    ]
    return doc


_CONSTRUCTORES = {
    "ventas": _doc_ventas,
    "productos-mas-vendidos": _doc_top,
    "inventario": _doc_inventario,
    "devoluciones": _doc_devoluciones,
    "rendimiento-vendedores": _doc_rendimiento,
}


def construir_documento(
    tipo: str, dto: Any, categoria: Optional[str] = None, extras: Optional[dict] = None
) -> Documento:
    """DTO del reporte -> Documento neutro (fuente unica de PDF y Excel)."""
    return _CONSTRUCTORES[tipo](dto, categoria, extras or {})


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
_NUM_FMT = {"moneda": '"Bs. "#,##0.00', "entero": "#,##0", "decimal": "#,##0.00"}
_BORDE = Border(bottom=Side(style="thin", color="D9E2F1"))
_RELLENO_CAB = PatternFill("solid", start_color=COLOR_PRIMARIO)
_RELLENO_SUAVE = PatternFill("solid", start_color=COLOR_SUAVE)
_RELLENO_AVISO = PatternFill("solid", start_color="EAF3FF")


def _escribir(ws, fila: int, col: int, valor: Any, tipo: TipoColumna = "texto"):
    """Escribe una celda. Texto siempre como cadena (sin formulas); numeros como numeros."""
    celda = ws.cell(row=fila, column=col)
    if valor is None:
        celda.value = "N/D"
        celda.data_type = "s"
        celda.alignment = Alignment(horizontal="right")
    elif tipo == "texto":
        celda.value = str(valor)
        celda.data_type = "s"  # nunca "f": evita inyeccion de formulas
    else:
        celda.value = float(valor) if tipo != "entero" else int(valor)
        celda.number_format = _NUM_FMT[tipo]
        celda.alignment = Alignment(horizontal="right")
    return celda


def _nombre_hoja(titulo: str, usados: set[str]) -> str:
    base = "".join(c for c in titulo if c not in "[]:*?/\\")[:31].strip() or "Hoja"
    nombre, n = base, 2
    while nombre.lower() in usados:
        sufijo = f" ({n})"
        nombre = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(nombre.lower())
    return nombre


def _xlsx(doc: Documento, generado_en: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    usados = {"resumen"}

    ws["A1"] = doc.titulo
    ws["A1"].font = Font(bold=True, size=16, color=COLOR_PRIMARIO)
    ws["A2"] = f"Generado: {generado_en.astimezone(timezone.utc):%Y-%m-%d %H:%M} (UTC) · Attention · CU20"
    ws["A2"].font = Font(italic=True, color="666666")

    fila = 4
    ws.cell(row=fila, column=1, value="Filtros aplicados").font = Font(bold=True, color=COLOR_PRIMARIO)
    fila += 1
    for etiqueta, valor in doc.filtros:
        c = ws.cell(row=fila, column=1, value=etiqueta)
        c.font = Font(bold=True)
        _escribir(ws, fila, 2, valor)
        fila += 1

    if doc.sin_datos:
        fila += 1
        c = _escribir(ws, fila, 1, doc.mensaje or MENSAJE_SIN_DATOS)
        c.font = Font(bold=True, color=COLOR_PRIMARIO)
        c.fill = _RELLENO_AVISO
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    else:
        fila += 1
        ws.cell(row=fila, column=1, value="Indicadores").font = Font(bold=True, color=COLOR_PRIMARIO)
        fila += 1
        for ind in doc.indicadores:
            ws.cell(row=fila, column=1, value=ind.etiqueta).font = Font(bold=True)
            _escribir(ws, fila, 2, ind.valor, ind.tipo)
            if ind.nota:
                _escribir(ws, fila, 3, ind.nota).font = Font(italic=True, color="666666")
            fila += 1
        for nota in doc.notas:
            fila += 1
            c = _escribir(ws, fila, 1, nota)
            c.font = Font(italic=True, color="666666")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
            ws.row_dimensions[fila].height = 32
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 32

    for sec in doc.secciones:
        hoja = wb.create_sheet(_nombre_hoja(sec.hoja or sec.titulo, usados))
        for j, col in enumerate(sec.columnas, start=1):
            c = hoja.cell(row=1, column=j, value=col.titulo)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _RELLENO_CAB
            c.alignment = Alignment(horizontal="right" if col.tipo != "texto" else "left", vertical="center")
        for i, valores in enumerate(sec.filas, start=2):
            for j, (col, valor) in enumerate(zip(sec.columnas, valores), start=1):
                _escribir(hoja, i, j, valor, col.tipo).border = _BORDE
        for j, col in enumerate(sec.columnas, start=1):
            ancho = max([len(col.titulo)] + [len(_texto_valor(f[j - 1], col.tipo)) for f in sec.filas]) + 3
            hoja.column_dimensions[get_column_letter(j)].width = min(max(ancho, 10), 60)
        hoja.freeze_panes = "A2"
        if sec.filas:
            hoja.auto_filter.ref = f"A1:{get_column_letter(len(sec.columnas))}{len(sec.filas) + 1}"
        if sec.nota:
            _escribir(hoja, len(sec.filas) + 3, 1, sec.nota).font = Font(italic=True, color="9A6700")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _pdf_texto(valor: Any) -> str:
    """Escapa marcado y sustituye lo que las fuentes base (WinAnsi) no dibujan."""
    limpio = str(valor).encode("cp1252", "replace").decode("cp1252")
    return escape(limpio)


class _LienzoNumerado(rl_canvas.Canvas):
    """Canvas con pie 'Página X de Y' (dos pasadas)."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._paginas: list[dict] = []

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for estado in self._paginas:
            self.__dict__.update(estado)
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#666666"))
            self.drawRightString(landscape(A4)[0] - 15 * mm, 8 * mm, f"Página {self._pageNumber} de {total}")
            self.drawString(15 * mm, 8 * mm, "Attention · CU20 · Gestión de Reportes")
            super().showPage()
        super().save()


def _pdf(doc: Documento, generado_en: datetime) -> bytes:
    base = getSampleStyleSheet()
    primario = colors.HexColor("#" + COLOR_PRIMARIO)
    est_titulo = ParagraphStyle("t", parent=base["Title"], fontName="Helvetica-Bold", fontSize=18, textColor=primario, alignment=TA_LEFT, spaceAfter=2)
    est_sub = ParagraphStyle("s", parent=base["Normal"], fontSize=8.5, textColor=colors.HexColor("#666666"), spaceAfter=8)
    est_h = ParagraphStyle("h", parent=base["Heading3"], fontName="Helvetica-Bold", fontSize=11, textColor=primario, spaceBefore=10, spaceAfter=4, keepWithNext=1)
    est_n = ParagraphStyle("n", parent=base["Normal"], fontSize=8.5, leading=11)
    est_nota = ParagraphStyle("nota", parent=est_n, textColor=colors.HexColor("#666666"), spaceBefore=4)
    est_celda = ParagraphStyle("c", parent=est_n, fontSize=8)
    est_celda_r = ParagraphStyle("cr", parent=est_celda, alignment=TA_RIGHT)

    buf = BytesIO()
    pagina = landscape(A4)
    ancho = pagina[0] - 30 * mm
    plantilla = SimpleDocTemplate(
        buf, pagesize=pagina, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
        title=doc.titulo, author="Attention", subject="CU20 - Gestión de Reportes",
        pageCompression=1 if COMPRIMIR_PDF else 0,
    )
    el: list[Any] = [
        Paragraph(_pdf_texto(doc.titulo), est_titulo),
        Paragraph(f"Generado: {generado_en.astimezone(timezone.utc):%Y-%m-%d %H:%M} (UTC)", est_sub),
    ]

    # Filtros aplicados
    el.append(Paragraph("Filtros aplicados", est_h))
    t = Table([[Paragraph(f"<b>{_pdf_texto(e)}</b>", est_celda), Paragraph(_pdf_texto(v), est_celda)] for e, v in doc.filtros],
              colWidths=[55 * mm, ancho - 55 * mm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + COLOR_SUAVE)),
                           ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F1")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    el.append(t)

    if doc.sin_datos:
        el.append(Spacer(1, 10))
        aviso = Table([[Paragraph(f"<b>{_pdf_texto(doc.mensaje or MENSAJE_SIN_DATOS)}</b>", est_n)]], colWidths=[ancho])
        aviso.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAF3FF")),
                                   ("BOX", (0, 0), (-1, -1), 0.5, primario), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        el.append(aviso)
    else:
        # Indicadores en filas de hasta 6 tarjetas
        el.append(Paragraph("Indicadores", est_h))
        por_fila = 6 if len(doc.indicadores) > 4 else 4
        celdas = [Paragraph(f"<font size=7 color='#1E4D8C'>{_pdf_texto(i.etiqueta.upper())}</font><br/>"
                            f"<font size=13><b>{_pdf_texto(_texto_valor(i.valor, i.tipo))}</b></font>"
                            + (f"<br/><font size=6.5 color='#666666'>{_pdf_texto(i.nota)}</font>" if i.nota else ""), est_n)
                  for i in doc.indicadores]
        filas = [celdas[k:k + por_fila] + [""] * (por_fila - len(celdas[k:k + por_fila])) for k in range(0, len(celdas), por_fila)]
        ti = Table(filas, colWidths=[ancho / por_fila] * por_fila)
        ti.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + COLOR_SUAVE)),
                                ("BOX", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F1")),
                                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.white), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        el.append(ti)

        for sec in doc.secciones:
            el.append(Paragraph(_pdf_texto(sec.titulo), est_h))
            if not sec.filas:
                el.append(Paragraph("Sin registros.", est_nota))
                continue
            cab = [Paragraph(f"<font color='white'><b>{_pdf_texto(c.titulo)}</b></font>", est_celda_r if c.tipo != "texto" else est_celda) for c in sec.columnas]
            cuerpo = [[Paragraph(_pdf_texto(_texto_valor(v, c.tipo)), est_celda if c.tipo == "texto" else est_celda_r)
                       for c, v in zip(sec.columnas, fila)] for fila in sec.filas]
            pesos = [max(len(c.titulo), min(max(len(_texto_valor(f[j], c.tipo)) for f in sec.filas), 38), 6) for j, c in enumerate(sec.columnas)]
            total = sum(pesos)
            tabla = Table([cab] + cuerpo, colWidths=[ancho * p / total for p in pesos], repeatRows=1)
            tabla.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), primario), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                       ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFF")]),
                                       ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F1")),
                                       ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            el.append(tabla)
            if sec.nota:
                el.append(Paragraph(_pdf_texto(sec.nota), est_nota))
        for nota in doc.notas:
            el.append(Paragraph(_pdf_texto(nota), est_nota))

    plantilla.build(el, canvasmaker=_LienzoNumerado)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# API publica
# ---------------------------------------------------------------------------
def exportar(
    tipo: str,
    formato: str,
    dto: Any,
    *,
    categoria: Optional[str] = None,
    extras: Optional[dict] = None,
    generado_en: Optional[datetime] = None,
) -> ArchivoExportado:
    """Genera el archivo del reporte a partir de su DTO.

    - `categoria`: nombre de la categoria filtrada (el DTO solo trae el id).
    - `extras`: parametros de la consulta que el DTO no repite (nivel_stock,
      top, limite_detalle, tipo_venta) para mostrarlos como filtros aplicados.
    """
    if formato not in MEDIA_TYPE:
        raise ValueError(f"formato no soportado: {formato}")
    generado = generado_en or datetime.now(timezone.utc)
    doc = construir_documento(tipo, dto, categoria, extras)
    contenido = _xlsx(doc, generado) if formato == "xlsx" else _pdf(doc, generado)
    ini, fin = doc.periodo
    return ArchivoExportado(contenido, MEDIA_TYPE[formato], formato, f"reporte-{tipo}_{ini}_{fin}.{formato}")
