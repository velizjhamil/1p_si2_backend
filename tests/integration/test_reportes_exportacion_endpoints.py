# Integracion CU20 - EXPORTACION PDF/Excel sobre la BD LOCAL de pruebas con el seed
# determinista. Solo corre contra `tienda_ropa_test` (candados de local_db).
#   $env:TEST_DATABASE_URL = "postgresql://postgres:CLAVE@localhost:5432/tienda_ropa_test"
#   python tests/integration/test_reportes_exportacion_endpoints.py
#
# Verifica: los 5 reportes en xlsx y pdf; que el archivo coincide con el JSON del
# mismo request (mismos filtros = mismos datos) y con las cifras calculadas a mano
# del seed; filtros aplicados rotulados; sin datos; 401/403/422 (JWT reales);
# lectura pura (sesion READ ONLY + conteos).
import re
import sys
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from tests.support import local_db  # noqa: E402

local_db.configure()  # ANTES de importar `app`

import app.main  # noqa: F401,E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from reportlab import rl_config  # noqa: E402
from sqlalchemy import text  # noqa: E402

from app.api.deps import get_current_user, get_db  # noqa: E402
from app.core.database import SessionLocal  # noqa: E402
from app.core.security import crear_token_acceso  # noqa: E402
from app.main import app  # noqa: E402
from app.modules.reportes import exportacion  # noqa: E402

BASE = "/api/v1/reportes"
MENSAJE = "No se encontraron datos para los parámetros ingresados."
TODO = {"fecha_inicio": "2025-12-01", "fecha_fin": "2026-03-31"}
SIN = {"fecha_inicio": "2020-01-01", "fecha_fin": "2020-01-10"}
REND = {"fecha_desde": "2025-12-01", "fecha_hasta": "2026-03-31"}
REND_SIN = {"fecha_desde": "2020-01-01", "fecha_hasta": "2020-01-10"}
REPORTES = {
    "ventas": TODO, "productos-mas-vendidos": TODO, "inventario": TODO,
    "devoluciones": TODO, "rendimiento-vendedores": REND,
}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _seed_disponible() -> bool:
    if local_db.modo_supabase_readonly() or not local_db.es_bd_de_pruebas_local():
        return False
    try:
        with SessionLocal() as db:
            return db.execute(text("select count(*) from ventas where codigo like 'ATT-T%'")).scalar() == 10
    except Exception:
        return False


def _usuario(rol):
    return SimpleNamespace(rol=SimpleNamespace(nombre_rol=rol))


def _pdf_legible(contenido: bytes) -> str:
    """Texto del PDF ya descomprimido (deshace escapes del formato)."""
    t = contenido.decode("latin-1")
    t = re.sub(r"\\([0-7]{3})", lambda m: chr(int(m.group(1), 8)), t)
    t = t.replace("\\(", "(").replace("\\)", ")")
    return re.sub(r"\)\s*Tj\s*\(", "", t)  # une los fragmentos de una misma línea (p. ej. por "&lt;")


@unittest.skipUnless(_seed_disponible(), "BD local con seed CU20 no disponible")
class Base(unittest.TestCase):
    def setUp(self):
        self.db = SessionLocal()
        self.db.execute(text("SET TRANSACTION READ ONLY"))
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: _usuario("GS")
        self.client = TestClient(app)
        self.cat = {n: i for i, n in self.db.execute(text("select id_categoria, nombre from categorias"))}

    def tearDown(self):
        app.dependency_overrides.clear()
        self.db.rollback()
        self.db.close()

    def get(self, ep, rol="GS", **params):
        app.dependency_overrides[get_current_user] = lambda: _usuario(rol)
        return self.client.get(f"{BASE}/{ep}", params=params)

    def archivo(self, ep, formato, params=None, **extra):
        r = self.get(ep, formato=formato, **(params if params is not None else REPORTES[ep]), **extra)
        self.assertEqual(r.status_code, 200, r.text[:300])
        return r

    def libro(self, ep, params=None, **extra):
        return load_workbook(BytesIO(self.archivo(ep, "xlsx", params, **extra).content))

    def pdf(self, ep, params=None, **extra):
        """PDF del endpoint con texto legible (sin compresión ni ASCII85) para las aserciones."""
        with mock.patch.object(exportacion, "COMPRIMIR_PDF", False), mock.patch.object(rl_config, "useA85", 0):
            return _pdf_legible(self.archivo(ep, "pdf", params, **extra).content)

    def json(self, ep, params=None, **extra):
        return self.get(ep, **(params if params is not None else REPORTES[ep]), **extra).json()["data"]


def filas(ws):
    return [tuple(f) for f in ws.iter_rows(values_only=True)]


def kv(ws):
    return {f[0]: f[1] for f in ws.iter_rows(values_only=True) if f[0] and f[1] is not None}


# ---------------------------------------------------------------------------
class ArchivosTest(Base):
    def test_los_5_reportes_en_xlsx_y_pdf(self):
        for ep in REPORTES:
            for fmt, tipo in (("xlsx", XLSX), ("pdf", "application/pdf")):
                r = self.archivo(ep, fmt)
                self.assertEqual(r.headers["content-type"], tipo, (ep, fmt))
                self.assertRegex(r.headers["content-disposition"], rf'^attachment; filename="reporte-{ep}_[0-9-]+_[0-9-]+\.{fmt}"$')
                self.assertEqual(r.headers["cache-control"], "no-store")
                self.assertGreater(len(r.content), 1000)
                self.assertTrue(r.content.startswith(b"PK" if fmt == "xlsx" else b"%PDF-"), (ep, fmt))
                if fmt == "pdf":
                    self.assertIn(b"%%EOF", r.content[-1024:])

    def test_nombre_del_archivo_refleja_el_periodo_aplicado(self):
        r = self.archivo("ventas", "xlsx", {"fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"})
        self.assertIn('filename="reporte-ventas_2026-01-01_2026-01-31.xlsx"', r.headers["content-disposition"])

    def test_formato_json_y_por_defecto_no_cambian(self):
        por_defecto = self.get("ventas", **TODO)
        explicito = self.get("ventas", formato="json", **TODO)
        self.assertEqual(por_defecto.json(), explicito.json())
        self.assertEqual(por_defecto.headers["content-type"], "application/json")
        self.assertEqual(por_defecto.json()["data"]["num_ventas"], 8)

    def test_pdf_grande_es_valido_con_compresion_de_produccion(self):
        for ep in REPORTES:
            r = self.archivo(ep, "pdf")  # sin patch: exactamente lo que se entrega
            self.assertTrue(r.content.startswith(b"%PDF-1."))


class ContenidoTest(Base):
    def test_ventas_xlsx_cifras_a_mano_y_coincide_con_json(self):
        wb = self.libro("ventas")
        r = kv(wb["Resumen"])
        self.assertEqual((r["Ingresos por productos"], r["Nº de ventas"], r["Unidades vendidas"], r["Ticket promedio"],
                          r["Total facturado"], r["Envíos cobrados"]), (2200, 8, 20, 275, 2260, 60))
        j = self.json("ventas")
        self.assertEqual(r["Ingresos por productos"], j["ingresos_productos"])
        self.assertEqual(filas(wb["Por categoría"])[1:], [("Camisas", 14, 1260), ("Pantalones", 5, 690), ("Vestidos", 1, 250)])
        self.assertEqual(filas(wb["Por canal"])[1:], [("POS (mostrador)", 4, 1150), ("Online", 4, 1050)])
        self.assertEqual(filas(wb["Por método de pago"])[1:], [("QR", 3, 1010), ("TARJETA", 2, 650), ("EFECTIVO", 3, 540)])
        self.assertEqual(filas(wb["Por día"])[1:], [(p["fecha"], p["num_ventas"], p["unidades"], p["ingresos"])
                                                    for p in j["por_fecha"] if p["num_ventas"] > 0])
        self.assertEqual(len(filas(wb["Por día"])) - 1, 7)

    def test_top_xlsx(self):
        wb = self.libro("productos-mas-vendidos")
        self.assertEqual(filas(wb["Ranking por unidades vendidas"])[1:],
                         [(1, "Camisa Oxford", "Camisas", 9, 900, 5), (2, "Jean Slim", "Pantalones", 3, 450, 2),
                          (3, "Camisa Lino", "Camisas", 3, 240, 2), (4, "Pantalon Chino", "Pantalones", 2, 240, 2),
                          (5, "Camisa Rayas", "Camisas", 2, 120, 1), (6, "Vestido Midi", "Vestidos", 1, 250, 1)])

    def test_inventario_xlsx_niveles_rotacion_y_nd(self):
        wb = self.libro("inventario")
        r = kv(wb["Resumen"])
        self.assertEqual((r["Productos"], r["Unidades en stock"], r["Valor del inventario"], r["Agotados"]), (8, 109, 13160, 1))
        self.assertEqual(filas(wb["Productos por nivel de stock"])[1:], [("Crítico (< 5)", 3), ("Bajo (< 15)", 2), ("OK", 3)])
        det = filas(wb["Detalle de inventario"])[1:]
        self.assertEqual([(d[0], d[3], d[4], d[8]) for d in det],
                         [("Jean Slim", 0, "CRITICO", "N/D"), ("Pantalon Chino", 3, "CRITICO", 0.67), ("Camisa Rayas", 4, "CRITICO", 0.5),
                          ("Vestido Midi", 5, "BAJO", 0.2), ("Camisa Lino", 12, "BAJO", 0.25), ("Vestido Gala", 15, "OK", 0),
                          ("Vestido Floral", 20, "OK", 0), ("Camisa Oxford", 50, "OK", 0.18)])  # stock 0 -> N/D, sin dividir
        self.assertNotIn("Chaqueta Inactiva", [d[0] for d in det])

    def test_devoluciones_xlsx(self):
        wb = self.libro("devoluciones")
        r = kv(wb["Resumen"])
        self.assertEqual((r["Devoluciones"], r["Unidades devueltas"], r["Importe devuelto"], r["Importe completado"]), (7, 7, 730, 550))
        self.assertEqual(filas(wb["Por estado"])[1:], [("APROBADA", 1, 1, 80), ("COMPLETADA", 4, 5, 550),
                                                       ("RECHAZADA", 1, 1, 100), ("SOLICITADA", 1, 1, 100)])
        self.assertEqual(filas(wb["Por producto"])[1:], [("Camisa Oxford", "Camisas", 3, 300), ("Camisa Lino", "Camisas", 2, 160),
                                                         ("Jean Slim", "Pantalones", 1, 150), ("Pantalon Chino", "Pantalones", 1, 120)])
        self.assertEqual(filas(wb["Por categoría"])[1:], [("Camisas", 5, 460), ("Pantalones", 2, 270)])
        det = filas(wb["Detalle de devoluciones"])[1:]
        self.assertEqual((len(det), det[0][1:3], det[3][3:]), (7, ("2026-03-12 10:00", "ATT-T00010"), ("RECHAZADA", "Cambie de opinion", 1, 100)))

    def test_rendimiento_xlsx(self):
        wb = self.libro("rendimiento-vendedores")
        fila = filas(wb["Rendimiento por vendedor"])[1]
        self.assertEqual((fila[0], fila[2], fila[3], fila[4], fila[5], fila[6]),
                         ("Vera Vendedora", 4, 1150.0, 287.5, "POS: 4", "2025-12-31 23:59"))
        self.assertEqual(kv(wb["Resumen"])["Ingresos"], 1150.0)

    def test_pdf_contiene_las_mismas_cifras(self):
        t = self.pdf("ventas")
        for x in ("Reporte de ventas", "Bs. 2,200.00", "Bs. 2,260.00", "Bs. 275.00", "Camisas", "Bs. 1,260.00", "POS (mostrador)", "2026-03-15"):
            self.assertIn(x, t, x)
        t = self.pdf("inventario")
        for x in ("Bs. 13,160.00", "Jean Slim", "N/D", "0.67", "Crítico (< 5)"):
            self.assertIn(x, t, x)
        self.assertNotIn("Chaqueta Inactiva", t)
        t = self.pdf("devoluciones")
        for x in ("Bs. 730.00", "Bs. 550.00", "ATT-T00010", "Costura rota", "12:00" if False else "2026-03-12 10:00"):
            self.assertIn(x, t, x)


class FiltrosAplicadosTest(Base):
    def test_filtros_del_reporte_ventas_xlsx_y_pdf(self):
        params = {**TODO, "categoria_id": self.cat["Pantalones"], "canal_venta": "pos"}
        wb = self.libro("ventas", params)
        f = kv(wb["Resumen"])
        self.assertEqual((f["Período (UTC)"], f["Categoría"], f["Canal de venta"]), ("2025-12-01 a 2026-03-31", "Pantalones", "POS (mostrador)"))
        self.assertEqual(f["Ingresos por productos"], 690)  # = pantalones + POS (ver test de datos)
        self.assertNotIn("Total facturado", f)
        j = self.json("ventas", params)
        self.assertEqual((f["Ingresos por productos"], f["Nº de ventas"]), (j["ingresos_productos"], j["num_ventas"]))
        t = self.pdf("ventas", params)
        for x in ("Pantalones", "POS (mostrador)", "Bs. 690.00", "no es atribuible"):
            self.assertIn(x, t, x)

    def test_el_archivo_respeta_cada_filtro_igual_que_el_json(self):
        casos = [
            ("productos-mas-vendidos", {**TODO, "top": 2}, "Ranking por unidades vendidas", 2),
            ("productos-mas-vendidos", {**TODO, "canal_venta": "POS"}, "Ranking por unidades vendidas", 4),
            ("inventario", {**TODO, "nivel_stock": "bajo"}, "Detalle de inventario", 2),
            ("inventario", {**TODO, "limite": 3}, "Detalle de inventario", 3),
            ("inventario", {**TODO, "categoria_id": self.cat["Vestidos"]}, "Detalle de inventario", 3),
            ("devoluciones", {**TODO, "estado": "completada"}, "Detalle de devoluciones", 4),
            ("devoluciones", {**TODO, "limite_detalle": 2}, "Detalle de devoluciones", 2),
            ("devoluciones", {**TODO, "canal_venta": "POS"}, "Detalle de devoluciones", 3),
        ]
        for ep, params, hoja, esperadas in casos:
            datos = [f for f in filas(self.libro(ep, params)[hoja])[1:] if f[0] is not None and not str(f[0]).startswith("Mostrando")]
            self.assertEqual(len(datos), esperadas, (ep, params))  # sin contar la nota bajo la tabla

    def test_filtros_extra_rotulados_en_el_archivo(self):
        f = kv(self.libro("inventario", {**TODO, "nivel_stock": "bajo", "limite": 5})["Resumen"])
        self.assertEqual((f["Nivel de stock"], f["Máx. de filas"], f["Ventana de rotación (UTC)"]), ("BAJO", "5", "2025-12-01 a 2026-03-31"))
        f = kv(self.libro("devoluciones", {**TODO, "estado": "RECHAZADA", "top": 3, "limite_detalle": 9})["Resumen"])
        self.assertEqual((f["Estado"], f["Top de productos"], f["Máx. filas del detalle"]), ("RECHAZADA", "3", "9"))
        f = kv(self.libro("rendimiento-vendedores", {**REND, "tipo_venta": "POS"})["Resumen"])
        self.assertEqual(f["Canal de venta"], "POS (mostrador)")
        self.assertEqual(kv(self.libro("productos-mas-vendidos", {**TODO, "top": 2})["Resumen"])["Top"], "2")

    def test_aviso_de_limite_igual_al_de_pantalla(self):
        t = " ".join(str(c) for f in filas(self.libro("inventario", {**TODO, "limite": 3})["Detalle de inventario"]) for c in f if c)
        self.assertIn("Mostrando 3 de 8", t)

    def test_cruce_de_anio_y_mes_en_la_exportacion(self):
        wb = self.libro("ventas", {"fecha_inicio": "2025-12-31", "fecha_fin": "2026-01-01"})
        self.assertEqual(kv(wb["Resumen"])["Ingresos por productos"], 770)
        wb = self.libro("ventas", {"fecha_inicio": "2026-02-28", "fecha_fin": "2026-03-01"})
        self.assertEqual(kv(wb["Resumen"])["Ingresos por productos"], 410)


class SinDatosTest(Base):
    PARAMS = {"ventas": SIN, "productos-mas-vendidos": SIN, "inventario": None, "devoluciones": SIN, "rendimiento-vendedores": REND_SIN}

    def test_sin_datos_genera_archivo_valido_con_el_mensaje(self):
        casos = dict(self.PARAMS)
        casos["inventario"] = {**TODO, "categoria_id": self.cat["Accesorios"]}  # categoría sin productos
        for ep, params in casos.items():
            wb = self.libro(ep, params)
            self.assertEqual(wb.sheetnames, ["Resumen"], ep)
            textos = [c for f in filas(wb["Resumen"]) for c in f if c]
            self.assertIn(MENSAJE, textos, ep)
            self.assertNotIn("Indicadores", textos, ep)
            t = self.pdf(ep, params)
            self.assertIn("No se encontraron datos para los par", t, ep)
            self.assertIn("Filtros aplicados", t, ep)

    def test_sin_datos_es_200_no_error(self):
        for ep, params in self.PARAMS.items():
            if params:
                self.assertEqual(self.get(ep, formato="xlsx", **params).status_code, 200, ep)
                self.assertEqual(self.get(ep, formato="pdf", **params).status_code, 200, ep)


class ErroresYPermisosTest(Base):
    def test_422_por_filtros_invalidos_tambien_al_exportar(self):
        for ep in ("ventas", "productos-mas-vendidos", "inventario", "devoluciones"):
            for fmt in ("xlsx", "pdf"):
                for params in ({"fecha_inicio": "2026-10-01", "fecha_fin": "2026-09-01"}, {"categoria_id": 99999}, {"canal_venta": "TIENDA"}):
                    r = self.get(ep, formato=fmt, **params)
                    self.assertEqual(r.status_code, 422, (ep, fmt, params))
                    self.assertEqual(r.headers["content-type"], "application/json")  # error legible, no un archivo roto
        self.assertEqual(self.get("rendimiento-vendedores", formato="xlsx", fecha_desde="2026-10-01", fecha_hasta="2026-09-01").status_code, 422)

    def test_422_formato_desconocido(self):
        for ep in REPORTES:
            for fmt in ("docx", "csv", "PDF ", ""):
                self.assertEqual(self.get(ep, formato=fmt).status_code, 422, (ep, fmt))

    def test_permisos_asu_gs_si_v_c_d_no(self):
        for ep in REPORTES:
            for fmt in ("xlsx", "pdf"):
                for rol in ("ASU", "GS"):
                    self.assertEqual(self.get(ep, rol=rol, formato=fmt).status_code, 200, (ep, fmt, rol))
                for rol in ("V", "C", "D"):
                    r = self.get(ep, rol=rol, formato=fmt)
                    self.assertEqual(r.status_code, 403, (ep, fmt, rol))
                    self.assertEqual(r.json()["detail"], "Su rol no tiene acceso al panel de reportes.")

    def test_403_antes_que_validar_el_formato_o_los_filtros(self):
        for rol in ("V", "C", "D"):
            for ep in REPORTES:
                self.assertEqual(self.get(ep, rol=rol, formato="docx", categoria_id=0).status_code, 403, (ep, rol))

    def test_403_no_genera_ningun_archivo(self):
        with mock.patch.object(exportacion, "exportar") as exportar:
            for rol in ("V", "C", "D"):
                for ep in REPORTES:
                    self.get(ep, rol=rol, formato="xlsx")
            exportar.assert_not_called()

    def test_exportar_no_modifica_datos(self):
        sql = ("select (select count(*) from ventas), (select count(*) from detalle_ventas), (select sum(stock_total) from productos), "
               "(select count(*) from devoluciones), (select count(*) from movimientos_inventario), (select count(*) from usuarios), "
               "(select max(ultima_conexion) from usuarios)")
        antes = self.db.execute(text(sql)).all()
        for ep in REPORTES:
            for fmt in ("xlsx", "pdf"):
                self.archivo(ep, fmt)
        self.assertEqual(antes, self.db.execute(text(sql)).all())


class JwtRealTest(Base):
    """Cadena de autenticación real (sin sobrescribir get_current_user): 401 / 403 / 200."""

    def setUp(self):
        super().setUp()
        app.dependency_overrides.pop(get_current_user, None)
        self.usuarios = {r: u for r, u in self.db.execute(text(
            "select distinct on (r.nombre_rol) r.nombre_rol, u.id_usuario from usuarios u join roles r on r.id_rol=u.id_rol where u.estado order by r.nombre_rol"))}

    def con_token(self, ep, rol, fmt):
        h = {"Authorization": "Bearer " + crear_token_acceso({"sub": str(self.usuarios[rol]), "rol": rol})}
        return self.client.get(f"{BASE}/{ep}", params={**REPORTES[ep], "formato": fmt}, headers=h)

    def test_sin_token_401_en_todos(self):
        for ep in REPORTES:
            for fmt in ("xlsx", "pdf"):
                r = self.client.get(f"{BASE}/{ep}", params={"formato": fmt})
                self.assertEqual(r.status_code, 401, (ep, fmt))
                self.assertEqual(r.headers.get("www-authenticate"), "Bearer")
        self.assertEqual(self.client.get(f"{BASE}/ventas", params={"formato": "docx"}).status_code, 401)  # 401 antes que 422

    def test_token_invalido_401(self):
        r = self.client.get(f"{BASE}/ventas", params={"formato": "xlsx"}, headers={"Authorization": "Bearer basura"})
        self.assertEqual(r.status_code, 401)

    def test_roles_con_jwt_reales(self):
        for ep in REPORTES:
            for fmt in ("xlsx", "pdf"):
                for rol in ("ASU", "GS"):
                    r = self.con_token(ep, rol, fmt)
                    self.assertEqual(r.status_code, 200, (ep, fmt, rol))
                    self.assertGreater(len(r.content), 1000)
                for rol in ("V", "C", "D"):
                    self.assertEqual(self.con_token(ep, rol, fmt).status_code, 403, (ep, fmt, rol))


if __name__ == "__main__":
    unittest.main()
