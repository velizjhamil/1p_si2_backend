# Pruebas unitarias AISLADAS de la exportacion CU20 (sin BD, sin red).
# Construyen DTOs a mano, generan XLSX/PDF y los leen de vuelta. Ejecutar
# desde la raiz del backend:
#   python tests/unit/test_reportes_exportacion.py
import re
import sys
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

# Red de seguridad: si algo intentara abrir la BD, que sea la local de pruebas, nunca Supabase.
from tests.support import local_db  # noqa: E402

local_db.configure()

from openpyxl import load_workbook  # noqa: E402

from app.modules.reportes import exportacion  # noqa: E402
from app.schemas.reporte import (  # noqa: E402
    MENSAJE_SIN_DATOS,
    DevolucionesReporteResponse,
    InventarioSituacionResponse,
    ProductosMasVendidosResponse,
    RendimientoVendedoresResponse,
    VentasPeriodoResponse,
)

FILTROS = {"fecha_inicio": date(2025, 12, 1), "fecha_fin": date(2026, 3, 31), "categoria_id": None, "canal_venta": None}
GENERADO = datetime(2026, 9, 20, 12, 0, tzinfo=timezone.utc)


def ventas(**extra):
    base = dict(
        filtros=FILTROS, sin_datos=False, num_ventas=8, unidades_vendidas=20,
        ingresos_productos=Decimal("2200"), ticket_promedio=Decimal("275"),
        total_facturado=Decimal("2260"), total_envios=Decimal("60"),
        por_fecha=[
            {"fecha": date(2025, 12, 30), "ingresos": Decimal("280"), "unidades": 3, "num_ventas": 1},
            {"fecha": date(2025, 12, 31), "ingresos": Decimal("0"), "unidades": 0, "num_ventas": 0},
        ],
        por_categoria=[{"id_categoria": 9, "categoria": "Camisas", "ingresos": Decimal("1260"), "unidades": 14}],
        por_canal=[{"canal": "POS", "ingresos": Decimal("1150"), "num_ventas": 4}],
        por_metodo_pago=[{"metodo_pago": "QR", "ingresos": Decimal("1010"), "num_ventas": 3}],
    )
    base.update(extra)
    return VentasPeriodoResponse.model_validate(base)


def ventas_vacio(**extra):
    return ventas(sin_datos=True, num_ventas=0, unidades_vendidas=0, ingresos_productos=Decimal("0"),
                  ticket_promedio=Decimal("0"), por_fecha=[], por_categoria=[], por_canal=[], por_metodo_pago=[], **extra)


def inventario(**extra):
    base = dict(
        filtros=FILTROS, sin_datos=False, total_productos=2, stock_total_unidades=50, valor_inventario=Decimal("5000"),
        agotados=1, por_nivel={"CRITICO": 1, "BAJO": 0, "OK": 1}, total_filas=5, limite=2,
        items=[
            {"id_producto": 1, "producto": "=HYPERLINK(\"http://x\")", "id_categoria": 1, "categoria": "C", "estado": "Agotado",
             "stock_actual": 0, "precio_venta": Decimal("150"), "valor_stock": Decimal("0"), "nivel_stock": "CRITICO",
             "unidades_vendidas": 3, "rotacion": None, "rotacion_disponible": False},
            {"id_producto": 2, "producto": "Camisa Ñandú <b>&</b>", "id_categoria": 1, "categoria": "C", "estado": "Activo",
             "stock_actual": 50, "precio_venta": Decimal("100"), "valor_stock": Decimal("5000"), "nivel_stock": "OK",
             "unidades_vendidas": 9, "rotacion": Decimal("0.18"), "rotacion_disponible": True},
        ],
    )
    base.update(extra)
    return InventarioSituacionResponse.model_validate(base)


def devoluciones():
    return DevolucionesReporteResponse.model_validate(dict(
        filtros=FILTROS, sin_datos=False, estado="COMPLETADA", num_devoluciones=4, unidades_devueltas=5,
        importe_total=Decimal("550"), importe_completado=Decimal("550"),
        por_estado=[{"estado": "COMPLETADA", "num_devoluciones": 4, "unidades": 5, "importe": Decimal("550")}],
        por_fecha=[{"fecha": date(2026, 3, 12), "num_devoluciones": 1, "unidades": 1, "importe": Decimal("120")}],
        por_producto=[{"id_producto": 1, "producto": "Camisa", "categoria": "Camisas", "unidades": 3, "importe": Decimal("300")}],
        por_categoria=[{"id_categoria": 1, "categoria": "Camisas", "unidades": 5, "importe": Decimal("460")}],
        detalle=[{"id_devolucion": 7, "fecha_solicitud": datetime(2026, 3, 12, 10, 0, tzinfo=timezone.utc), "estado": "COMPLETADA",
                  "codigo_venta": "ATT-T00010", "motivo": "+cmd|' /C calc'!A0", "unidades": 1, "importe": Decimal("120")}],
    ))


def rendimiento(vacio=False):
    items = [] if vacio else [{
        "id_vendedor": "abc", "nombre": "Vera", "correo": "v@x.com", "total_ventas": 4, "total_ingresos": Decimal("1150.00"),
        "ticket_promedio": Decimal("287.50"), "primera_venta": datetime(2025, 12, 31, 23, 59, 59, tzinfo=timezone.utc),
        "ultima_venta": None, "tipos_venta": {"POS": 4}}]
    return RendimientoVendedoresResponse(items=items, total_vendedores=len(items), total_ingresos=Decimal("0" if vacio else "1150.00"),
                                         total_operaciones=0 if vacio else 4, fecha_desde=date(2025, 12, 1), fecha_hasta=date(2026, 3, 31))


def xlsx(tipo, dto, **kw):
    a = exportacion.exportar(tipo, "xlsx", dto, generado_en=GENERADO, **kw)
    return a, load_workbook(BytesIO(a.contenido))


def pdf_texto(tipo, dto, **kw):
    """PDF sin compresion: el texto queda legible en el binario para las pruebas."""
    from reportlab import rl_config

    anterior, exportacion.COMPRIMIR_PDF = exportacion.COMPRIMIR_PDF, False
    a85, rl_config.useA85 = rl_config.useA85, 0  # sin ASCII85: texto legible en el binario
    try:
        a = exportacion.exportar(tipo, "pdf", dto, generado_en=GENERADO, **kw)
    finally:
        exportacion.COMPRIMIR_PDF, rl_config.useA85 = anterior, a85
    crudo = a.contenido.decode("latin-1")
    # deshace los escapes del formato PDF: barra+paréntesis y octales (barra+341 = á)
    texto = re.sub(r"\\([0-7]{3})", lambda m: chr(int(m.group(1), 8)), crudo)
    texto = texto.replace("\\(", "(").replace("\\)", ")")
    return a, re.sub(r"\)\s*Tj\s*\(", "", texto)  # une fragmentos de una misma línea


class ExcelTest(unittest.TestCase):
    def test_hojas_valores_y_tipos_numericos(self):
        a, wb = xlsx("ventas", ventas(), categoria=None)
        self.assertEqual(wb.sheetnames, ["Resumen", "Por día", "Por categoría", "Por canal", "Por método de pago"])
        r = {fila[0]: fila[1] for fila in wb["Resumen"].iter_rows(values_only=True) if fila[0]}
        self.assertEqual((r["Ingresos por productos"], r["Nº de ventas"], r["Ticket promedio"], r["Total facturado"]), (2200, 8, 275, 2260))
        celda = next(c for fila in wb["Resumen"].iter_rows() for c in fila if c.value == "Ingresos por productos")
        valor = wb["Resumen"].cell(row=celda.row, column=2)
        self.assertIsInstance(valor.value, (int, float))  # número, no texto
        self.assertIn("Bs.", valor.number_format)
        dia = list(wb["Por día"].iter_rows(values_only=True))
        self.assertEqual(dia, [("Fecha", "Ventas", "Unidades", "Ingresos"), ("2025-12-30", 1, 3, 280)])  # solo días con ventas
        self.assertEqual(a.nombre, "reporte-ventas_2025-12-01_2026-03-31.xlsx")
        self.assertEqual(a.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_filtros_aplicados_en_el_resumen(self):
        f = dict(FILTROS, categoria_id=9, canal_venta="POS")
        _, wb = xlsx("ventas", ventas(filtros=f, total_facturado=None, total_envios=None), categoria="Camisas")
        filas = {x[0]: x[1] for x in wb["Resumen"].iter_rows(values_only=True) if x[0] and x[1]}
        self.assertEqual((filas["Período (UTC)"], filas["Categoría"], filas["Canal de venta"]),
                         ("2025-12-01 a 2026-03-31", "Camisas", "POS (mostrador)"))
        self.assertNotIn("Total facturado", filas)  # con categoría no se informa
        textos = " ".join(str(c) for x in wb["Resumen"].iter_rows(values_only=True) for c in x if c)
        self.assertIn("no es atribuible", textos)

    def test_sin_datos_una_sola_hoja_con_el_mensaje(self):
        _, wb = xlsx("ventas", ventas_vacio(mensaje=MENSAJE_SIN_DATOS))
        self.assertEqual(wb.sheetnames, ["Resumen"])
        textos = [c for x in wb["Resumen"].iter_rows(values_only=True) for c in x if c]
        self.assertIn(MENSAJE_SIN_DATOS, textos)
        self.assertNotIn("Indicadores", textos)

    def test_no_evalua_formulas_ni_inyeccion(self):
        _, wb = xlsx("inventario", inventario(), extras={"nivel_stock": "CRITICO"})
        celda = wb["Detalle de inventario"]["A2"]
        self.assertEqual(celda.value, '=HYPERLINK("http://x")')
        self.assertEqual(celda.data_type, "s")  # cadena, NO fórmula
        _, wb = xlsx("devoluciones", devoluciones())
        motivo = [c for c in list(wb["Detalle de devoluciones"].iter_rows(min_row=2, values_only=False))[0] if str(c.value).startswith("+cmd")][0]
        self.assertEqual(motivo.data_type, "s")

    def test_rotacion_no_calculable_es_nd_y_extras_como_filtros(self):
        _, wb = xlsx("inventario", inventario(), extras={"nivel_stock": "CRITICO"})
        det = list(wb["Detalle de inventario"].iter_rows(min_row=2, values_only=True))
        self.assertEqual((det[0][8], det[1][8]), ("N/D", 0.18))
        filtros = {x[0]: x[1] for x in wb["Resumen"].iter_rows(values_only=True) if x[0] and x[1]}
        self.assertEqual((filtros["Nivel de stock"], filtros["Máx. de filas"]), ("CRITICO", "2"))
        textos = " ".join(str(c) for x in wb["Detalle de inventario"].iter_rows(values_only=True) for c in x if c)
        self.assertIn("Mostrando 2 de 5", textos)  # aviso de límite igual al de pantalla

    def test_devoluciones_y_rendimiento(self):
        _, wb = xlsx("devoluciones", devoluciones(), extras={"top": 3, "limite_detalle": 20})
        f = {x[0]: x[1] for x in wb["Resumen"].iter_rows(values_only=True) if x[0] and x[1]}
        self.assertEqual((f["Estado"], f["Top de productos"], f["Máx. filas del detalle"]), ("COMPLETADA", "3", "20"))
        det = list(wb["Detalle de devoluciones"].iter_rows(min_row=2, values_only=True))[0]
        self.assertEqual((det[1], det[6]), ("2026-03-12 10:00", 120))  # UTC
        _, wb = xlsx("rendimiento-vendedores", rendimiento(), extras={"tipo_venta": "POS"})
        fila = list(wb["Rendimiento por vendedor"].iter_rows(min_row=2, values_only=True))[0]
        self.assertEqual((fila[0], fila[2], fila[3], fila[4], fila[7]), ("Vera", 4, 1150.0, 287.5, "—"))  # string decimal -> número
        f = {x[0]: x[1] for x in wb["Resumen"].iter_rows(values_only=True) if x[0] and x[1]}
        self.assertEqual(f["Canal de venta"], "POS (mostrador)")

    def test_rendimiento_sin_vendedores_es_sin_datos(self):
        _, wb = xlsx("rendimiento-vendedores", rendimiento(vacio=True))
        self.assertEqual(wb.sheetnames, ["Resumen"])
        self.assertIn(MENSAJE_SIN_DATOS, [c for x in wb["Resumen"].iter_rows(values_only=True) for c in x if c])


class PdfTest(unittest.TestCase):
    def test_pdf_valido_con_datos(self):
        a = exportacion.exportar("ventas", "pdf", ventas(), generado_en=GENERADO)  # con compresión (producción)
        self.assertTrue(a.contenido.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", a.contenido[-1024:])
        self.assertEqual((a.media_type, a.nombre), ("application/pdf", "reporte-ventas_2025-12-01_2026-03-31.pdf"))

    def test_contenido_filtros_y_kpis(self):
        f = dict(FILTROS, categoria_id=9, canal_venta="POS")
        _, t = pdf_texto("ventas", ventas(filtros=f), categoria="Camisas")
        for esperado in ("Reporte de ventas", "2025-12-01 a 2026-03-31", "Camisas", "POS (mostrador)", "Bs. 2,200.00",
                         "Bs. 2,260.00", "2025-12-30", "P\\341gina 1 de 1", "Generado: 2026-09-20 12:00 (UTC)"):
            self.assertIn(esperado.replace("\\341", "\xe1"), t, esperado)
        self.assertNotIn("2025-12-31", t)  # día sin ventas no se lista

    def test_sin_datos_mensaje_y_sin_tablas(self):
        _, t = pdf_texto("ventas", ventas_vacio(mensaje=MENSAJE_SIN_DATOS))
        self.assertIn("No se encontraron datos para los par", t)
        self.assertNotIn("Indicadores", t)
        _, t = pdf_texto("rendimiento-vendedores", rendimiento(vacio=True))
        self.assertIn("No se encontraron datos", t)

    def test_texto_hostil_no_rompe_el_pdf(self):
        a, t = pdf_texto("inventario", inventario())
        self.assertTrue(a.contenido.startswith(b"%PDF-"))
        self.assertIn("N/D", t)  # rotación no calculable
        self.assertIn("Camisa Ñandú <b>&</b>", t)  # el marcado se escapó: se ve como texto literal, no como negrita
        raro = inventario(items=[dict(inventario().items[0].model_dump(), producto="日本語 ✓ producto")])
        self.assertTrue(exportacion.exportar("inventario", "pdf", raro).contenido.startswith(b"%PDF-"))  # no lanza

    def test_tabla_larga_pagina_con_encabezado_repetido(self):
        muchos = [dict(inventario().items[1].model_dump(), id_producto=i, producto=f"Producto {i}") for i in range(120)]
        _, t = pdf_texto("inventario", inventario(items=muchos, total_filas=120, limite=200))
        self.assertGreaterEqual(t.count("Producto"), 121)  # 120 filas + encabezado repetido
        self.assertRegex(t, r"Página 1 de [2-9]")  # varias páginas


class ApiTest(unittest.TestCase):
    def test_formato_invalido_y_tipos(self):
        with self.assertRaises(ValueError):
            exportacion.exportar("ventas", "docx", ventas())
        for tipo, dto in (("productos-mas-vendidos", ProductosMasVendidosResponse.model_validate(dict(
                filtros=FILTROS, sin_datos=False, top=1, total_productos_vendidos=1,
                items=[{"posicion": 1, "id_producto": 1, "producto": "P", "id_categoria": 1, "categoria": "C",
                        "cantidad_vendida": 9, "total_generado": Decimal("900"), "num_ventas": 5}]))),):
            for fmt in ("xlsx", "pdf"):
                self.assertGreater(len(exportacion.exportar(tipo, fmt, dto).contenido), 500)

    def test_no_hay_acceso_a_datos(self):
        # La exportación es pura: no importa SQLAlchemy ni sesiones.
        fuente = Path(exportacion.__file__).read_text(encoding="utf-8")
        for prohibido in ("sqlalchemy", "Session", "SessionLocal", "db.query"):
            self.assertNotIn(prohibido, fuente)


if __name__ == "__main__":
    unittest.main()
